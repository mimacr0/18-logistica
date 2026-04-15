# -*- coding: utf-8 -*-
##############################################################################
#
# Copyright 2026 DaFe Solutions
#
##############################################################################

import base64
import csv
import io
import logging
from collections import Counter
from datetime import date, datetime

from odoo import models, fields, api, _
from odoo.exceptions import UserError, ValidationError

_logger = logging.getLogger(__name__)

# Columnas aceptadas (cabecera CSV, sin distinguir mayúsculas)
CSV_ALIASES = {
    'account_sku': {'account_sku', 'sku', 'client_sku', 'referencia_cliente'},
    'product': {'product', 'product_code', 'default_code', 'internal_ref', 'codigo_interno'},
    'tracking': {'tracking', 'trazabilidad', 'trace'},
    'account_ean13': {'account_ean13', 'ean13', 'ean'},
    'account_asin': {'account_asin', 'asin'},
    'account_fnsku': {'account_fnsku', 'fnsku'},
    'notes': {'notes', 'notas', 'note'},
}


class AccountProductMapImportWizard(models.TransientModel):
    _name = 'account.product.map.import.wizard'
    _description = 'Import Product Map from CSV or Excel'

    account_id = fields.Many2one('account.partner', string='Client Account', required=True)
    file_data = fields.Binary(string='CSV or Excel file', required=True)
    filename = fields.Char(string='Filename')
    dry_run = fields.Boolean(
        string='Dry run only',
        default=True,
        help='If enabled, validates the file and shows a report without creating or updating records.',
    )
    delimiter = fields.Char(string='Delimiter', size=1, default=',')
    has_header = fields.Boolean(string='File has header row', default=True)
    result_html = fields.Html(string='Result', readonly=True, sanitize=False)
    is_excel_file = fields.Boolean(
        string='Excel file',
        compute='_compute_is_excel_file',
        help='True when the filename is .xlsx or .xls (delimiter applies only to CSV).',
    )

    @api.depends('filename', 'file_data')
    def _compute_is_excel_file(self):
        for wiz in self:
            wiz.is_excel_file = wiz._spreadsheet_kind() in ('xlsx', 'xls')

    def _spreadsheet_kind(self):
        """Return 'csv', 'xlsx' or 'xls' (uses filename extension and optional binary sniff)."""
        self.ensure_one()
        fn = (self.filename or '').lower()
        if fn.endswith('.xlsx'):
            return 'xlsx'
        if fn.endswith('.xls'):
            return 'xls'
        if self.file_data:
            head = base64.b64decode(self.file_data)[:8]
            if len(head) >= 2 and head[:2] == b'PK':
                return 'xlsx'
            ole2 = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'
            if len(head) >= 8 and head[:8] == ole2:
                return 'xls'
        return 'csv'

    @api.model
    def _cell_to_text(self, val):
        if val is None:
            return ''
        if isinstance(val, bool):
            return '1' if val else '0'
        if isinstance(val, (int,)):
            return str(val)
        if isinstance(val, float):
            if val == int(val):
                return str(int(val))
            return str(val).strip()
        if isinstance(val, datetime):
            return val.strftime('%Y-%m-%d %H:%M:%S')
        if isinstance(val, date):
            return val.strftime('%Y-%m-%d')
        return str(val).strip()

    def _read_xlsx_rows_raw(self, raw):
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise UserError(
                _('The Python package openpyxl is required to import .xlsx files. (%s)') % str(e)
            ) from e
        wb = load_workbook(filename=io.BytesIO(raw), read_only=True, data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            rows = []
            for row in ws.iter_rows(values_only=True):
                if row is None:
                    continue
                line = [self._cell_to_text(c) for c in row]
                rows.append(line)
            while rows and not any(rows[-1]):
                rows.pop()
        finally:
            wb.close()
        return rows

    def _read_xls_rows_raw(self, raw):
        try:
            import xlrd
        except ImportError as e:
            raise UserError(
                _('The Python package xlrd is required to import .xls files. (%s)') % str(e)
            ) from e
        book = xlrd.open_workbook(file_contents=raw)
        sheet = book.sheet_by_index(0)
        rows = []
        for rx in range(sheet.nrows):
            line = []
            for cx in range(sheet.ncols):
                cell = sheet.cell(rx, cx)
                if cell.ctype == xlrd.XL_CELL_EMPTY:
                    line.append('')
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    v = cell.value
                    line.append(str(int(v)) if v == int(v) else str(v))
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        dt = xlrd.xldate_as_datetime(cell.value, book.datemode)
                        line.append(dt.strftime('%Y-%m-%d %H:%M:%S'))
                    except Exception:
                        line.append(str(cell.value))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    line.append('1' if cell.value else '0')
                else:
                    line.append(str(cell.value).strip())
            rows.append(line)
        while rows and not any(rows[-1]):
            rows.pop()
        return rows

    def _rows_to_parsed(self, rows):
        """Build (row_index, line_dict) list from a list of cell string lists."""
        self.ensure_one()
        if not rows:
            raise UserError(_('The file is empty.'))
        start = 1 if self.has_header else 0
        if self.has_header:
            header = [self._normalize_header(c) for c in rows[0]]
            idx_map = {}
            for i, h in enumerate(header):
                if not h:
                    continue
                idx_map[i] = h
            data_rows = rows[1:]
        else:
            idx_map = dict(enumerate([
                'account_sku', 'product', 'tracking',
                'account_ean13', 'account_asin', 'account_fnsku', 'notes',
            ]))
            data_rows = rows

        parsed = []
        for row_idx, row in enumerate(data_rows, start=start + 1):
            line = {}
            for col_idx, cell in enumerate(row):
                if col_idx not in idx_map:
                    continue
                canon = idx_map[col_idx]
                line[canon] = str(cell).strip() if cell is not None else ''
            parsed.append((row_idx, line))
        return parsed

    def _normalize_header(self, key):
        if key is None:
            return None
        k = key.strip().lower()
        k = k.replace(' ', '_').replace('-', '_')
        for canonical, aliases in CSV_ALIASES.items():
            if k in aliases:
                return canonical
        return k

    def _parse_selection_value(self, field_name, raw, Model):
        field = Model._fields[field_name]
        if not raw or not str(raw).strip():
            return False
        val = str(raw).strip().lower()
        selection = field._description_selection(self.env)
        keys = [k for k, _lbl in selection if k is not False and k is not None]
        if val in keys:
            return val
        for key, label in selection:
            if key is False or key is None:
                continue
            if label and str(label).lower() == val:
                return key
        allowed = ', '.join(keys)
        raise ValueError(_('Invalid %(field)s: %(value)s (allowed: %(allowed)s)') % {
            'field': field_name,
            'value': raw,
            'allowed': allowed,
        })

    def _find_product_by_ref(self, code):
        code = (code or '').strip()
        if not code:
            return self.env['product.product']
        Product = self.env['product.product'].sudo()
        product = Product.search([('default_code', '=', code)], limit=1)
        if product:
            return product
        product = Product.search([('barcode', '=', code)], limit=1)
        return product

    def _read_csv_rows_raw(self, raw):
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('latin-1')
        sep = (self.delimiter or ',')[:1] or ','
        reader = csv.reader(io.StringIO(text), delimiter=sep)
        rows = list(reader)
        if not rows:
            raise UserError(_('The file is empty.'))
        return rows

    def _read_data_rows(self):
        """Decode uploaded file to list of rows (list of string cells)."""
        self.ensure_one()
        if not self.file_data:
            raise UserError(_('Please upload a file.'))
        raw = base64.b64decode(self.file_data)
        fmt = self._spreadsheet_kind()
        if fmt == 'xlsx':
            return self._read_xlsx_rows_raw(raw)
        if fmt == 'xls':
            return self._read_xls_rows_raw(raw)
        return self._read_csv_rows_raw(raw)

    def _read_csv_rows(self):
        """Return parsed (row_no, line_dict) tuples (CSV, XLS or XLSX)."""
        rows = self._read_data_rows()
        return self._rows_to_parsed(rows)

    def _build_report_html(self, lines):
        rows_html = []
        for row_no, status, msg, extra in lines:
            color = '#c62828' if status == 'error' else ('#1565c0' if status == 'skip' else '#2e7d32')
            rows_html.append(
                '<tr><td>%s</td><td style="color:%s"><strong>%s</strong></td><td>%s</td><td>%s</td></tr>'
                % (row_no, color, status, msg, extra or '')
            )
        return (
            '<table class="table table-sm table-bordered o_list_view">'
            '<thead><tr><th>Row</th><th>Status</th><th>Message</th><th>Detail</th></tr></thead>'
            '<tbody>%s</tbody></table>'
        ) % (''.join(rows_html))

    def action_run(self):
        self.ensure_one()
        Map = self.env['account.product.map'].sudo()
        report_lines = []
        parsed = self._read_csv_rows()
        sku_in_file = Counter()
        for row_idx, line in parsed:
            sku = (line.get('account_sku') or '').strip()
            if sku:
                sku_in_file[sku] += 1

        created = updated = errors = 0
        for row_idx, line in parsed:
            sku = (line.get('account_sku') or '').strip()
            product_ref = (line.get('product') or '').strip()
            extra_detail = ''

            def add_line(status, msg, detail=''):
                nonlocal errors
                report_lines.append((row_idx, status, msg, detail or extra_detail))
                if status == 'error':
                    errors += 1

            try:
                if not sku:
                    add_line('error', _('Missing account_sku'))
                    continue
                if sku_in_file[sku] > 1:
                    add_line('error', _('Duplicate account_sku in file'))
                    continue
                if not product_ref:
                    add_line('error', _('Missing product (internal code / barcode)'))
                    continue

                product = self._find_product_by_ref(product_ref)
                if not product:
                    add_line('error', _('No internal product found for reference'), product_ref)
                    continue

                vals = {
                    'account_id': self.account_id.id,
                    'account_sku': sku,
                    'product_id': product.id,
                }
                if line.get('tracking'):
                    vals['tracking'] = self._parse_selection_value('tracking', line['tracking'], Map)
                for opt in ('account_ean13', 'account_asin', 'account_fnsku', 'notes'):
                    if line.get(opt):
                        vals[opt] = line[opt].strip()

                existing = Map.search([
                    ('account_id', '=', self.account_id.id),
                    ('account_sku', '=', sku),
                ], limit=1)

                if self.dry_run:
                    if existing:
                        add_line('ok', _('Would update existing mapping'), existing.display_name)
                    else:
                        add_line('ok', _('Would create new mapping'), product.display_name)
                    continue

                if existing:
                    existing.write(vals)
                    updated += 1
                    add_line('ok', _('Updated'), existing.display_name)
                else:
                    Map.create(vals)
                    created += 1
                    add_line('ok', _('Created'), product.display_name)

            except ValueError as e:
                add_line('error', str(e))
            except ValidationError as e:
                add_line('error', str(e))
            except Exception as e:
                _logger.exception('Import row %s failed', row_idx)
                add_line('error', _('Error: %s') % str(e))

        summary = _(
            '<p><strong>Summary:</strong> created %(c)d, updated %(u)d, rows with errors %(e)d. Dry run: %(dry)s</p>'
        ) % {'c': created, 'u': updated, 'e': errors, 'dry': _('Yes') if self.dry_run else _('No')}

        self.result_html = summary + self._build_report_html(report_lines)

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'account.product.map.import.wizard',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
            'context': self.env.context,
        }
